from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\beauty")
SOURCE_DOCX = ROOT / "毕业论文_最终稿_规范达标版.docx"
WORK_DOCX = ROOT / "tmp" / "thesis_restructured_working.docx"
ASSET_DIR = ROOT / "final_assets" / "thesis_restructure"
SCREENSHOT_DIR = ROOT / "screenshots" / "thesis_pages"
BACKEND_DIR = ROOT / "beauty-knowledge-backend"
FRONTEND_DIR = ROOT / "beauty-knowledge-frontend"

FONT_SONG = Path(r"C:\Windows\Fonts\simsun.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\simhei.ttf")
FONT_MONO = Path(r"C:\Windows\Fonts\consola.ttf")


@dataclass
class AssetBlock:
    title: str
    path: Path
    intro: str
    analysis: str
    width_cm: float = 13.6


def get_font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size=size)


def style_run(run, font_name: str = "宋体", size_pt: float = 14, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    for run in paragraph.runs:
        style_run(run, "宋体", 14)


def style_heading_1(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.page_break_before = True
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(28)
    for run in paragraph.runs:
        style_run(run, "黑体", 16, bold=True)


def style_heading_2(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(28)
    for run in paragraph.runs:
        style_run(run, "黑体", 14, bold=True)


def style_heading_3(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(28)
    for run in paragraph.runs:
        style_run(run, "黑体", 14, bold=True)


def style_caption(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(28)
    for run in paragraph.runs:
        style_run(run, "宋体", 14)


def style_image_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.0


def set_table_font(table: Table, font_name: str = "宋体", size_pt: float = 10.5) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = Pt(22)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    style_run(run, font_name, size_pt)


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def insert_table_after(paragraph: Paragraph, rows: int, cols: int) -> Table:
    doc = paragraph._parent
    table = doc.add_table(rows=rows, cols=cols, width=Cm(15.5))
    paragraph._p.addnext(table._tbl)
    return table


def remove_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)
        paragraph._p = paragraph._element = None


def find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == exact_text:
            return para
    raise ValueError(f"Paragraph not found: {exact_text}")


def find_paragraph_contains(doc: Document, needle: str) -> Paragraph:
    for para in doc.paragraphs:
        if needle in para.text:
            return para
    raise ValueError(f"Paragraph not found containing: {needle}")


def paragraphs_between(doc: Document, start_text: str, end_text: str) -> list[Paragraph]:
    collecting = False
    items: list[Paragraph] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text == start_text:
            collecting = True
        if collecting:
            items.append(para)
        if text == end_text and collecting:
            break
    return items


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], width: int = 3) -> None:
    draw.line([start, end], fill="black", width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 14 * direction, ey - 8), (ex - 14 * direction, ey + 8)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 8, ey - 14 * direction), (ex + 8, ey - 14 * direction)]
    draw.polygon(points, fill="black")


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
                       font: ImageFont.FreeTypeFont) -> None:
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=6, align="center")
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    tx = x1 + (x2 - x1 - tw) / 2
    ty = y1 + (y2 - y1 - th) / 2
    draw.multiline_text((tx, ty), text, fill="black", font=font, align="center", spacing=6)


def draw_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str,
             font: ImageFont.FreeTypeFont, radius: int = 0) -> None:
    if radius:
        draw.rounded_rectangle(box, outline="black", width=3, radius=radius, fill="white")
    else:
        draw.rectangle(box, outline="black", width=3, fill="white")
    draw_centered_text(draw, box, text, font)


def create_data_flow_diagram(output: Path) -> None:
    img = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(FONT_BOLD, 38)
    body_font = get_font(FONT_SONG, 26)
    small_font = get_font(FONT_SONG, 22)

    draw.text((630, 35), "系统需求分析数据流图", fill="black", font=title_font)
    draw_box(draw, (90, 180, 330, 300), "普通用户", body_font, radius=24)
    draw_box(draw, (90, 640, 330, 760), "管理员", body_font, radius=24)
    draw_box(draw, (1420, 150, 1710, 290), "问答服务\n与会话管理", body_font, radius=18)
    draw_box(draw, (760, 170, 1060, 300), "登录鉴权\n与角色路由", body_font, radius=18)
    draw_box(draw, (520, 410, 800, 560), "知识检索\n与详情访问", body_font, radius=18)
    draw_box(draw, (860, 410, 1140, 560), "文件接入\n与任务处理", body_font, radius=18)
    draw_box(draw, (1210, 410, 1490, 560), "实体审核\n与图谱构建", body_font, radius=18)
    draw_box(draw, (540, 780, 820, 920), "MySQL\n业务数据", body_font, radius=18)
    draw_box(draw, (900, 780, 1180, 920), "MinIO\n原始文件", body_font, radius=18)
    draw_box(draw, (1260, 780, 1540, 920), "Milvus\n向量索引", body_font, radius=18)

    draw_arrow(draw, (330, 240), (760, 240))
    draw_arrow(draw, (330, 700), (760, 240))
    draw_arrow(draw, (1060, 240), (1420, 220))
    draw_arrow(draw, (290, 240), (520, 460))
    draw_arrow(draw, (290, 700), (860, 460))
    draw_arrow(draw, (800, 485), (1420, 220))
    draw_arrow(draw, (1140, 485), (1210, 485))
    draw_arrow(draw, (1210, 520), (1540, 850))
    draw_arrow(draw, (680, 560), (680, 780))
    draw_arrow(draw, (950, 560), (1040, 780))
    draw_arrow(draw, (1380, 560), (1380, 780))

    draw.text((355, 205), "登录请求 / 权限校验", fill="black", font=small_font)
    draw.text((365, 675), "资料接入 / 后台治理", fill="black", font=small_font)
    draw.text((1115, 180), "问答调用", fill="black", font=small_font)
    draw.text((338, 410), "知识浏览", fill="black", font=small_font)
    draw.text((1130, 595), "审核结果写入图谱", fill="black", font=small_font)

    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)


def create_module_diagram(output: Path) -> None:
    img = Image.new("RGB", (1800, 1000), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(FONT_BOLD, 38)
    body_font = get_font(FONT_SONG, 24)

    draw.text((620, 30), "系统功能模块设计图", fill="black", font=title_font)
    draw_box(draw, (720, 120, 1080, 220), "美业多媒体知识库系统", get_font(FONT_BOLD, 28), radius=20)

    boxes = [
        ((120, 320, 420, 460), "用户端模块\n首页 / 问答 / 收藏"),
        ((470, 320, 770, 460), "管理端模块\n总览 / 报表 / 用户 / 公告"),
        ((820, 320, 1120, 460), "知识治理模块\n知识 / 分类 / 详情"),
        ((1170, 320, 1470, 460), "文件处理模块\n上传 / 解析 / 任务"),
        ((310, 620, 610, 760), "实体确认模块\n候选审核 / 状态回写"),
        ((660, 620, 960, 760), "知识图谱模块\n图谱查询 / 路径 / 证据"),
        ((1010, 620, 1310, 760), "问答服务模块\n检索增强 / SSE / 会话"),
        ((1360, 620, 1660, 760), "基础设施模块\nMySQL / Redis / MQ / MinIO / Milvus"),
    ]
    for box, text in boxes:
        draw_box(draw, box, text, body_font, radius=18)
        draw_arrow(draw, (900, 220), ((box[0] + box[2]) // 2, box[1]))

    draw_arrow(draw, (610, 690), (660, 690))
    draw_arrow(draw, (960, 690), (1010, 690))
    draw_arrow(draw, (1310, 690), (1360, 690))

    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)


def create_flowchart(output: Path, title: str, steps: list[str]) -> None:
    img = Image.new("RGB", (1400, 1700), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(FONT_BOLD, 36)
    body_font = get_font(FONT_SONG, 24)
    draw.text((420, 35), title, fill="black", font=title_font)
    top = 140
    centers: list[tuple[int, int]] = []
    for idx, step in enumerate(steps):
        box = (330, top + idx * 210, 1070, top + idx * 210 + 120)
        draw_box(draw, box, step, body_font, radius=20)
        centers.append(((box[0] + box[2]) // 2, box[3]))
        if idx:
            prev = centers[idx - 1]
            draw_arrow(draw, (prev[0], prev[1] + 8), ((box[0] + box[2]) // 2, box[1] - 8))
    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)


def read_lines(path: Path, start: int, end: int) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    snippet = []
    for idx in range(start - 1, min(end, len(lines))):
        snippet.append(f"{idx + 1:>3}  {lines[idx].rstrip()}")
    return "\n".join(snippet)


def create_code_image(output: Path, title: str, code: str) -> None:
    title_font = get_font(FONT_BOLD, 26)
    code_font = get_font(FONT_MONO, 22)
    line_height = 34
    padding = 32
    title_h = 70
    lines = code.splitlines() or [""]
    width = 1700
    height = title_h + padding * 2 + len(lines) * line_height + 20
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, title_h), fill="#f5f5f5", outline="black", width=2)
    draw.text((28, 20), title, fill="black", font=title_font)
    y = title_h + padding
    for line in lines:
        draw.text((padding, y), line.replace("\t", "    "), fill="black", font=code_font)
        y += line_height
    draw.rectangle((1, 1, width - 2, height - 2), outline="black", width=2)
    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)


def generate_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = {
        "data_flow": ASSET_DIR / "fig_3_2_data_flow.png",
        "module_design": ASSET_DIR / "fig_4_2_module_design.png",
        "ingest_flow": ASSET_DIR / "fig_4_4_ingest_flow.png",
        "qa_flow": ASSET_DIR / "fig_4_5_qa_flow.png",
        "code_router": ASSET_DIR / "fig_5_21_code_router.png",
        "code_knowledge": ASSET_DIR / "fig_5_22_code_knowledge_controller.png",
        "code_file": ASSET_DIR / "fig_5_23_code_file_controller.png",
        "code_kg": ASSET_DIR / "fig_5_24_code_kg_controller.png",
        "code_chat": ASSET_DIR / "fig_5_25_code_chat_controller.png",
    }
    create_data_flow_diagram(assets["data_flow"])
    create_module_diagram(assets["module_design"])
    create_flowchart(
        assets["ingest_flow"],
        "知识入库处理流程图",
        ["管理员创建知识或上传原始文件", "系统完成文件存储并登记处理任务", "OCR/转写与文本清洗生成可处理文本",
         "切片、实体抽取与候选关系生成", "管理员在实体确认页面完成审核", "知识、图谱与向量索引同步更新"],
    )
    create_flowchart(
        assets["qa_flow"],
        "用户问答服务流程图",
        ["用户输入问题或上传附件摘要请求", "系统执行关键词检索与向量召回", "融合排序并组织证据片段",
         "构建会话上下文与提示词", "问答服务以流式方式返回答案", "前端展示来源、消息历史和附件关联信息"],
    )
    create_code_image(
        assets["code_router"],
        "前端统一路由与角色鉴权实现（router/index.ts）",
        read_lines(FRONTEND_DIR / "src" / "router" / "index.ts", 1, 62),
    )
    create_code_image(
        assets["code_knowledge"],
        "知识管理接口控制器实现（KnowledgeController.java）",
        read_lines(
            BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "cms" / "controller" / "KnowledgeController.java",
            1,
            67,
        ),
    )
    create_code_image(
        assets["code_file"],
        "文件上传与任务追踪接口实现（FileController.java）",
        read_lines(
            BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "cms" / "controller" / "FileController.java",
            1,
            95,
        ),
    )
    create_code_image(
        assets["code_kg"],
        "知识图谱查询与审核接口实现（KgController.java）",
        read_lines(
            BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "kg" / "controller" / "KgController.java",
            1,
            118,
        ),
    )
    create_code_image(
        assets["code_chat"],
        "智能问答流式响应接口实现（ChatController.java）",
        read_lines(
            BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "rag" / "controller" / "ChatController.java",
            1,
            132,
        ),
    )
    return assets


def add_bookmark(paragraph: Paragraph, bookmark_name: str, bookmark_id: int) -> None:
    if any(
        child.tag == qn("w:bookmarkStart") and child.get(qn("w:name")) == bookmark_name
        for child in paragraph._p.iterchildren()
    ):
        return
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    p.insert(0, start)
    p.append(end)


def append_internal_hyperlink(paragraph: Paragraph, text: str, anchor: str, superscript: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    if superscript:
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    rpr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "21")
    rpr.append(size_cs)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def rebuild_citations(paragraph: Paragraph) -> None:
    text = paragraph.text
    if not re.search(r"\[(\d+)\]", text):
        return
    p_pr = paragraph._p.pPr
    clear_paragraph(paragraph)
    if p_pr is not None and paragraph._p.pPr is None:
        paragraph._p.insert(0, p_pr)
    cursor = 0
    for match in re.finditer(r"\[(\d+)\]", text):
        normal = text[cursor:match.start()]
        if normal:
            run = paragraph.add_run(normal)
            style_run(run, "宋体", 10.5)
        ref_no = match.group(1)
        append_internal_hyperlink(paragraph, f"[{ref_no}]", f"gzcu_ref_{ref_no}", superscript=True)
        cursor = match.end()
    tail = text[cursor:]
    if tail:
        run = paragraph.add_run(tail)
        style_run(run, "宋体", 10.5)


def is_heading_text(text: str) -> bool:
    return bool(re.match(r"^(第[一二三四五六七八九十]+章|\d+\.\d+(\.\d+)?\s|结\s*论|参考文献|致\s*谢)", text))


def add_reference_hyperlinks(doc: Document) -> None:
    ref_section = False
    bookmark_id = 100
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt == "参考文献":
            ref_section = True
            continue
        if ref_section:
            if txt in {"致    谢", "致谢"}:
                break
            match = re.match(r"^\[(\d+)\]", txt)
            if match:
                add_bookmark(para, f"gzcu_ref_{match.group(1)}", bookmark_id)
                bookmark_id += 1

    body_started = False
    skip_summary = False
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        if txt == "第一章 绪论":
            body_started = True
        if txt in {"结    论", "结论"}:
            break
        if not body_started:
            continue
        if txt.endswith("本章小结"):
            skip_summary = True
            continue
        if skip_summary:
            if is_heading_text(txt):
                skip_summary = False
            else:
                continue
        if re.search(r"\[(\d+)\]", txt):
            rebuild_citations(para)


def insert_body_after(paragraph: Paragraph, text: str) -> Paragraph:
    para = insert_paragraph_after(paragraph, text)
    style_body_paragraph(para)
    return para


def insert_heading2_after(paragraph: Paragraph, text: str) -> Paragraph:
    para = insert_paragraph_after(paragraph, text)
    style_heading_2(para)
    return para


def insert_heading3_after(paragraph: Paragraph, text: str) -> Paragraph:
    para = insert_paragraph_after(paragraph, text)
    style_heading_3(para)
    return para


def insert_captioned_image(after: Paragraph, caption: str, image_path: Path, intro: str, analysis: str,
                           width_cm: float = 13.6) -> Paragraph:
    current = insert_body_after(after, intro)
    image_para = insert_paragraph_after(current)
    image_para.add_run().add_picture(str(image_path), width=Cm(width_cm))
    style_image_paragraph(image_para)
    caption_para = insert_paragraph_after(image_para, caption)
    style_caption(caption_para)
    current = insert_body_after(caption_para, analysis)
    return current


def build_interface_table(after: Paragraph, doc: Document) -> Paragraph:
    intro = insert_body_after(
        after,
        "结合系统真实控制器实现与前端页面调用关系，可将对论文论述最关键的接口整理为表 5-2。"
        "这些接口覆盖登录鉴权、知识分页、文件上传、任务追踪、实体审核、图谱查询和流式问答等核心链路，"
        "能够直接对应系统的主要业务能力。",
    )
    caption = insert_paragraph_after(intro, "表 5-2 系统核心接口设计表")
    style_caption(caption)
    table = insert_table_after(caption, rows=9, cols=7)
    headers = ["序号", "接口名称", "请求路径", "方法", "功能说明", "主要参数", "返回结果"]
    rows = [
        ["1", "用户登录", "/api/auth/login", "POST", "完成身份认证并返回 token 与角色信息", "username、password、role", "登录状态与用户信息"],
        ["2", "知识分页查询", "/api/knowledge/page", "GET", "分页检索知识条目并返回列表", "pageNum、pageSize、keyword、status", "知识分页数据"],
        ["3", "知识搜索", "/api/knowledge/search", "GET", "面向前台场景的全文检索", "keyword、pageNum、pageSize", "搜索结果列表"],
        ["4", "文件上传", "/api/file/upload", "POST", "上传原始资料并创建处理任务", "file、knowledgeId、categoryId、fileType", "文件与任务信息"],
        ["5", "任务查询", "/api/file/task/{taskId}", "GET", "查看文件处理进度和执行状态", "taskId", "任务详情与状态"],
        ["6", "待审核图谱候选", "/api/kg/pending", "GET", "查询实体抽取待确认结果", "status", "候选实体与关系列表"],
        ["7", "图谱查询", "/api/kg/product/{id}/graph", "GET", "按产品查询图谱节点和关系", "id", "图谱节点与边结果"],
        ["8", "问答流式接口", "/api/chat/stream", "GET/POST", "基于检索增强链路返回流式回答", "question、sessionId、knowledgeId", "SSE 分片结果"],
    ]
    for idx, head in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = head
        for run in cell.paragraphs[0].runs:
            style_run(run, "宋体", 10.5, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
            table.cell(r, c).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c in {0, 3} else WD_ALIGN_PARAGRAPH.LEFT
    set_table_font(table)
    note = insert_paragraph_after(caption)
    table._tbl.addnext(note._p)
    set_paragraph_text(note, "从表 5-2 可以看出，系统接口设计遵循资源化路径与角色分权控制原则，"
                             "前端每一类页面操作都能在后端找到清晰的映射接口，这也是系统详细设计能够落到真实实现的关键。")
    style_body_paragraph(note)
    return note


def screenshot_blocks() -> list[AssetBlock]:
    return [
        AssetBlock("图 5-7 系统登录页面截图", SCREENSHOT_DIR / "fig_5_3_login_page.png", "在系统详细设计阶段，登录页面是角色入口和权限分流的直接体现，因此首先展示系统登录界面的真实运行效果。", "图 5-7 中可以看到账号密码输入、角色提示和登录入口等元素，说明前端已经围绕统一认证入口完成了界面组织，并与后端登录鉴权接口形成了稳定映射。"),
        AssetBlock("图 5-8 普通用户端首页页面截图", SCREENSHOT_DIR / "fig_5_4_user_home_page.png", "用户端首页是知识服务能力的第一展示面，页面中集中体现了推荐内容、公告入口和快速访问区域。", "图 5-8 说明系统已经能够在同一页面中组织推荐知识、公告信息和问答入口，使用户在进入平台后可以直接完成知识浏览与进一步交互。"),
        AssetBlock("图 5-9 普通用户端智能问答页面截图", SCREENSHOT_DIR / "fig_5_5_user_chat_page.png", "智能问答页面对应系统详细设计中的问答模块，是检索增强与流式输出能力最直接的前端体现。", "图 5-9 展示了会话列表、消息区域、快捷提问和输入控件，说明系统在详细实现层已经形成了连续对话、消息回显和问题输入的完整界面结构。"),
        AssetBlock("图 5-10 普通用户端收藏页面截图", SCREENSHOT_DIR / "fig_5_6_user_favorites_page.png", "收藏页面用于承接用户对知识条目的二次使用需求，也是知识复用能力在前端的具体体现。", "图 5-10 显示收藏列表、检索与继续提问入口已经完成整合，表明系统能够将用户感兴趣的知识内容沉淀为可持续访问的个人记录。"),
        AssetBlock("图 5-11 管理员端运营总览页面截图", SCREENSHOT_DIR / "fig_5_7_admin_overview_page.png", "后台总览页面承接管理员掌握系统状态的需求，是管理端首页的重要组成部分。", "图 5-11 中的知识统计、任务状态和待确认实体等信息，体现出系统已经把平台运行指标组织为可观察的概览数据。"),
        AssetBlock("图 5-12 管理员端运营报表页面截图", SCREENSHOT_DIR / "fig_5_8_admin_report_page.png", "运营报表页面用于支持后台数据分析，也是系统面向治理侧扩展统计能力的重要界面。", "图 5-12 展示了活跃用户、热点内容和来源占比等统计信息，说明项目实现中已具备基础报表与运营分析能力。"),
        AssetBlock("图 5-13 管理员端知识列表页面截图", SCREENSHOT_DIR / "fig_5_9_admin_knowledge_list_page.png", "知识列表页面对应知识管理模块的核心前端实现，用于支撑知识的查询、筛选与维护。", "图 5-13 可以看到列表、筛选条件和状态信息已经联动展示，说明知识管理模块的页面与接口在实现层已经形成闭环。"),
        AssetBlock("图 5-14 管理员端文件上传页面截图", SCREENSHOT_DIR / "fig_5_10_admin_upload_page.png", "文件上传页面对应文件接入与任务处理模块，是原始资料进入系统的主要入口。", "图 5-14 中的上传控件、知识信息表单和提交按钮表明系统已经支持通过界面创建处理任务，并为后续 OCR、转写和实体抽取提供触发点。"),
        AssetBlock("图 5-15 管理员端分类树页面截图", SCREENSHOT_DIR / "fig_5_11_admin_category_page.png", "分类树页面体现了系统对知识分层组织的实现效果，能够直观展示分类节点与维护操作。", "图 5-15 说明系统已支持层级分类的浏览、查看和节点维护，这也是知识治理能力在详细设计层的重要组成部分。"),
        AssetBlock("图 5-16 管理员端实体确认页面截图", SCREENSHOT_DIR / "fig_5_12_admin_entity_confirm_page.png", "实体确认页面对应人工审核环节，是将自动抽取结果转化为正式图谱数据的关键界面。", "图 5-16 展示了待审核候选、统计信息与审核动作区，反映出系统在详细实现中保留了人工确认机制以控制知识质量。"),
        AssetBlock("图 5-17 管理员端知识图谱页面截图", SCREENSHOT_DIR / "fig_5_13_admin_graph_page.png", "知识图谱页面用于可视化展示实体与关系，是系统结构化知识组织能力的主要前端表现。", "图 5-17 中的图谱查询区和节点关系展示结果说明系统已经支持管理员从界面层直接查看图结构关系及其组织效果。"),
        AssetBlock("图 5-18 管理员端任务监控页面截图", SCREENSHOT_DIR / "fig_5_14_admin_task_monitor_page.png", "任务监控页面用于跟踪异步处理链路中的任务状态，是文件处理模块可观测性的核心界面。", "图 5-18 展示了任务状态、时间信息和重试操作，表明系统已经能够对资料处理任务进行持续跟踪与故障恢复。"),
        AssetBlock("图 5-19 管理员端用户管理页面截图", SCREENSHOT_DIR / "fig_5_15_admin_user_manage_page.png", "用户管理页面用于体现系统在权限控制与用户治理方面的实现效果。", "图 5-19 显示了用户列表、角色信息和状态切换入口，说明系统已经具备管理员侧的基础用户管理与角色控制能力。"),
        AssetBlock("图 5-20 管理员端公告管理页面截图", SCREENSHOT_DIR / "fig_5_16_admin_notice_page.png", "公告管理页面用于承接平台通知发布与维护功能，是管理端运营能力的组成部分。", "图 5-20 展示公告列表、发布状态与操作按钮，能够说明系统详细设计不仅覆盖知识治理，也兼顾了平台运营维护场景。"),
    ]


def code_blocks(assets: dict[str, Path]) -> list[AssetBlock]:
    return [
        AssetBlock("图 5-21 前端统一路由与角色鉴权代码截图", assets["code_router"], "图 5-21 选取前端路由定义与守卫代码，用于说明系统如何在页面层区分普通用户与管理员入口。", "代码中同时定义了 `/user` 与 `/admin` 两类路由树，并通过 `beforeEach` 守卫读取登录状态与角色信息，从而保证页面访问权限与论文前文所述角色边界保持一致。"),
        AssetBlock("图 5-22 知识管理接口控制器代码截图", assets["code_knowledge"], "图 5-22 展示知识管理控制器的主要实现，用于说明知识分页、详情、保存、更新和搜索能力的接口组织方式。", "可以看出该控制器围绕知识资源统一定义了分页、详情与状态更新接口，同时通过 `PreAuthorize` 对管理操作做权限约束，体现出知识治理与前台消费共用同一资源模型的实现思路。"),
        AssetBlock("图 5-23 文件上传与任务追踪接口代码截图", assets["code_file"], "图 5-23 选取文件处理控制器代码，展示文件上传、任务查询、重试和文件打开等关键接口实现。", "该代码能够直接说明文件处理模块并不是简单上传功能，而是围绕任务状态、对象访问和异常重试组织了一套完整的处理链路。"),
        AssetBlock("图 5-24 知识图谱接口控制器代码截图", assets["code_kg"], "图 5-24 选取图谱控制器代码，用于展示候选审核、图谱查询、路径分析和证据查看接口的组织方式。", "代码中既包含面向管理员的审核接口，也包含图谱查询和证据查询接口，说明图谱模块同时承担数据治理出口和结构化查询入口两类职责。"),
        AssetBlock("图 5-25 智能问答流式接口代码截图", assets["code_chat"], "图 5-25 展示智能问答控制器中流式问答、附件摘要、附件追问和会话管理相关代码。", "从实现细节可以看出，问答模块通过 SSE 返回流式结果，并围绕附件摘要、会话历史和消息列表组织接口，这与前文页面截图中的问答界面形成了直接对应关系。"),
    ]


def insert_chapter2(doc: Document) -> None:
    anchor = find_paragraph(doc, "本章围绕课题的研究背景、研究意义、国内外研究现状以及论文整体结构进行了阐述，明确了本文的研究对象、研究内容与论述思路，为后续系统需求分析和设计实现奠定了基础。")
    current = insert_paragraph_after(anchor, "第二章 相关技术与理论基础")
    style_heading_1(current)
    current = insert_heading2_after(current, "2.1 前后端分离与系统技术栈")
    current = insert_body_after(current, "本系统采用前后端分离架构组织实现。前端部分基于 Vue 3、Vue Router、Pinia 和 Element Plus 构建页面层，主要负责用户端与管理端界面的交互呈现、角色入口划分以及状态同步；后端部分基于 Spring Boot 3、MyBatis-Plus 与 Spring Security 组织业务服务、数据访问和权限控制逻辑。前后端通过 HTTP 接口协同，使页面交互、业务处理和数据存储边界保持清晰，有利于在毕业设计规模内完成模块化开发与后续扩展。")
    current = insert_body_after(current, "从项目代码可以看出，前端统一通过路由守卫控制管理员和普通用户的访问边界，后端则围绕认证授权、知识管理、文件处理、图谱服务和问答服务划分控制器与服务层。这种技术组织方式既适合构建以页面交互为主的应用系统，也能够为知识治理与智能服务并存的业务场景提供足够的工程支撑。")
    current = insert_heading2_after(current, "2.2 多媒体资料处理与知识接入基础")
    current = insert_body_after(current, "系统面对的知识来源并不限于结构化条目，还包括说明书、培训资料、图片和音频等半结构化或非结构化内容。因此，项目在文件接入层引入了 MinIO 对象存储、OCR 文本识别与 Whisper 语音转写能力，把原始文件逐步转化为可清洗、可切片、可抽取的文本结果。文档智能处理研究表明，对非结构化资料进行识别、解析和结构化组织是构建行业知识底座的重要前提[9][10]。")
    current = insert_body_after(current, "在工程实现上，系统并未将文件解析结果直接用于前台展示，而是先进入任务表、切片表与候选实体结果中间层，再由后台审核与图谱构建模块继续处理。这样既保留了处理链路的可追踪性，也使知识接入从一次性操作转变为可治理、可回溯的连续过程。")
    current = insert_heading2_after(current, "2.3 知识图谱与结构化知识组织基础")
    current = insert_body_after(current, "知识图谱为系统提供了从“文本集合”走向“结构化知识网络”的表达能力。在美业场景中，产品、成分、功效、适用人群与使用方式之间具有明显的语义关联，仅依赖普通文本检索难以完整呈现它们之间的关系。相关研究指出，知识图谱在推荐系统、问答系统和证据追溯场景中能够显著增强语义关联表达、关系发现和可解释性[6][11]。")
    current = insert_body_after(current, "考虑到自动抽取结果可能存在关系失真或错误传播风险，系统在图谱构建链路中保留了实体候选确认、关系审核和证据查看机制，以降低未经校验的数据直接进入正式图谱的风险[12]。这与面向真实场景的知识治理思路相一致，也为后续图谱查询和问答解释提供了更可靠的结构化基础。")
    current = insert_heading2_after(current, "2.4 检索增强生成与流式问答基础")
    current = insert_body_after(current, "系统智能问答模块采用检索增强生成思路，在回答生成前先完成关键词检索、向量召回、证据融合与上下文构建。相关研究表明，RAG 能够通过引入外部知识源降低模型幻觉风险，并提升回答的事实性与时效性[1][3]。当图结构知识与层级知识底座参与召回过程时，系统还能够进一步增强证据组织和答案解释能力[2][4][13]。")
    current = insert_body_after(current, "结合本项目代码实现，问答服务并不是孤立的大模型调用，而是建立在知识条目、文本切片、向量索引和会话上下文之上的组合式服务链路。同时，SSE 流式返回机制使前端可以逐步展示回答内容、来源信息和附件关联结果，从而让问答模块兼具交互性与可解释性。")
    current = insert_heading2_after(current, "2.5 本章小结")
    insert_body_after(current, "本章围绕系统真实采用的前后端技术栈、多媒体资料处理能力、知识图谱组织方式以及检索增强生成问答机制展开阐述，说明了系统能够完成知识治理与智能服务的技术基础，为后续需求分析、总体设计和详细实现章节提供了理论与工程依据。")


def update_existing_structure(doc: Document) -> None:
    exact_map = {
        "第二章 系统需求分析": "第三章 系统需求与分析", "2.1 系统设计目标": "3.1 系统设计目标", "2.2 功能需求分析": "3.2 功能需求分析",
        "2.3 用例图与关键用例描述": "3.4 用例图与关键用例描述", "2.3.1 普通用户用例图分析": "3.4.1 普通用户用例图分析", "2.3.2 后台管理员用例图分析": "3.4.2 后台管理员用例图分析",
        "2.3.3 关键用例描述表": "3.4.3 关键用例描述表", "2.4 非功能需求分析": "3.5 非功能需求分析", "2.5 可行性分析": "3.6 可行性分析", "2.6 本章小结": "3.7 本章小结",
        "图 2-1 系统业务流程图": "图 3-1 系统业务流程图", "图 2-2 普通用户用例图": "图 3-3 普通用户用例图", "图 2-3 后台管理员用例图": "图 3-4 后台管理员用例图",
        "表 2-1 智能问答用例描述表": "表 3-1 智能问答用例描述表", "表 2-2 文件处理与实体确认用例描述表": "表 3-2 文件处理与实体确认用例描述表",
        "第三章 系统总体设计": "第四章 系统总体设计", "3.1 系统整体结构设计": "4.1 系统整体结构设计", "3.2 功能模块设计": "4.2 功能模块设计",
        "3.3 数据库设计": "4.3 数据库设计", "3.4 主要数据表设计": "4.4 主要数据表设计", "3.5 关键流程设计": "4.5 关键流程设计", "3.6 本章小结": "4.6 本章小结",
        "图 3-1 系统整体结构图": "图 4-1 系统整体结构图", "图 3-2 系统 ER 图": "图 4-3 系统 ER 图", "表 3-1 主要业务数据表设计": "表 4-1 主要业务数据表设计", "表 3-2 图谱与问答相关数据表设计": "表 4-2 图谱与问答相关数据表设计",
        "第四章 系统详细设计与实现": "第五章 系统详细设计与实现", "4.1 设计思路概述": "5.1 设计思路概述", "4.2 前端交互模块设计与实现": "5.2 前端交互模块设计与实现", "4.3 知识管理模块设计与实现": "5.3 知识管理模块设计与实现",
        "4.4 文件处理模块设计与实现": "5.4 文件处理模块设计与实现", "4.5 实体抽取与知识图谱模块设计与实现": "5.5 实体抽取与知识图谱模块设计与实现", "4.6 智能问答模块设计与实现": "5.6 智能问答模块设计与实现",
        "4.7 接口设计说明": "5.7 接口设计说明", "4.8 本章小结": "5.9 本章小结", "图 4-1 系统详细设计模块协作图": "图 5-1 系统详细设计模块协作图", "图 4-2 用户端页面跳转关系图": "图 5-2 用户端页面跳转关系图",
        "图 4-3 知识管理模块处理流程图": "图 5-3 知识管理模块处理流程图", "图 4-4 文件处理模块流程图": "图 5-4 文件处理模块流程图", "图 4-5 知识图谱查询流程图": "图 5-5 知识图谱查询流程图", "图 4-6 智能问答模块流程图": "图 5-6 智能问答模块流程图",
        "表 4-1 统一响应格式说明表": "表 5-1 统一响应格式说明表", "第五章 系统测试": "第六章 系统测试", "5.1 测试目标与测试环境": "6.1 测试目标与测试环境", "5.2 功能测试": "6.2 功能测试", "5.3 性能测试与兼容性测试": "6.3 性能测试与兼容性测试", "5.4 安全性测试": "6.4 安全性测试", "5.5 测试结果分析": "6.5 测试结果分析", "5.6 本章小结": "6.6 本章小结", "图 5-1 性能测试记录图": "图 6-1 性能测试记录图", "图 5-2 兼容性测试结果图": "图 6-2 兼容性测试结果图",
        "表 5-2 用户登录功能测试用例": "表 6-2 用户登录功能测试用例", "表 5-3 知识管理功能测试用例": "表 6-3 知识管理功能测试用例",
        "表 5-4 文件处理功能测试用例": "表 6-4 文件处理功能测试用例", "表 5-5 知识图谱与问答功能测试用例": "表 6-5 知识图谱与问答功能测试用例",
        "表 5-6 安全性测试用例": "表 6-6 安全性测试用例",
        "第六章 系统部署与运行验证": "第七章 系统部署与运行验证", "6.1 部署方案设计": "7.1 部署方案设计", "6.2 运行链路验证": "7.2 运行链路验证", "6.3 工程化特点分析": "7.3 工程化特点分析", "6.4 本章小结": "7.4 本章小结",
    }
    replacements = [("第二章分析系统需求；第三章介绍系统总体设计；第四章详细说明系统各模块的设计与实现；第五章对系统测试情况进行分析；第六章介绍系统部署与运行验证；", "第二章介绍系统相关技术与理论基础；第三章分析系统需求与数据流；第四章介绍系统总体设计与关键流程；第五章详细说明系统各模块的设计实现、运行界面与关键代码；第六章对系统测试情况进行分析；第七章介绍系统部署与运行验证；"), ("图 2-1 所示", "图 3-1 所示"), ("图 2-2 与图 2-3", "图 3-3 与图 3-4"), ("如图 2-2 所示", "如图 3-3 所示"), ("如图 2-3 所示", "如图 3-4 所示"), ("系统整体结构如图 3-1 所示", "系统整体结构如图 4-1 所示"), ("系统主要实体及关系组织如图 3-2 所示", "系统主要实体及关系组织如图 4-3 所示"), ("模块协作关系如图 4-1 所示", "模块协作关系如图 5-1 所示"), ("主要页面之间的跳转关系如图 4-2 所示", "主要页面之间的跳转关系如图 5-2 所示"), ("处理流程如图 4-3 所示", "处理流程如图 5-3 所示"), ("核心流程如图 4-4 所示", "核心流程如图 5-4 所示"), ("如图 4-5 所示", "如图 5-5 所示"), ("如图 4-6 所示", "如图 5-6 所示"), ("统一响应结构如表 4-1 所示", "统一响应结构如表 5-1 所示"), ("图 5-1、图 5-2", "图 6-1、图 6-2"), ("表 5-1 可用于整理测试环境配置信息", "表 6-1 可用于整理测试环境配置信息")]
    for para in doc.paragraphs:
        text = para.text.strip()
        if text in exact_map:
            set_paragraph_text(para, exact_map[text])
        else:
            new_text = para.text
            for old, new in replacements:
                new_text = new_text.replace(old, new)
            if new_text != para.text:
                set_paragraph_text(para, new_text)


def insert_requirement_data_flow(doc: Document, assets: dict[str, Path]) -> None:
    caption = find_paragraph(doc, "图 3-1 系统业务流程图")
    current = insert_heading2_after(caption, "3.3 系统数据流分析")
    current = insert_body_after(current, "除业务流程外，需求分析阶段还需要进一步说明系统在真实运行中涉及的数据来源、处理节点和输出方向。本系统的数据流既包含普通用户发起的知识浏览与问答请求，也包含管理员发起的文件接入、知识治理和图谱审核流程。")
    current = insert_body_after(current, "如图 3-2 所示，系统在接收用户请求后，会经过登录鉴权、知识检索、文件处理、实体审核与问答服务等多个处理节点，同时分别与关系型数据库、对象存储和向量索引进行交互。这说明系统的需求不仅是页面功能堆叠，更是多类数据在治理与服务之间的有序流转。")
    insert_captioned_image(current, "图 3-2 系统数据流图", assets["data_flow"], "为更直观地呈现系统在需求层面的输入、处理和输出关系，本文对用户端、管理端与后台服务之间的数据流向进行了梳理。", "图中将普通用户、管理员、业务处理模块及三类核心数据存储进行统一表达，能够说明系统需求分析阶段已经考虑了资料接入、知识沉淀与问答输出之间的完整闭环。")


def insert_overall_design_assets(doc: Document, assets: dict[str, Path]) -> None:
    after_module = find_paragraph(doc, "从模块职责划分看，用户服务模块重点解决“用户如何高效使用知识”的问题，后台管理模块重点解决“管理员如何治理知识”的问题，知识处理模块重点解决“原始资料如何转化为结构化资源”的问题，问答服务模块重点解决“如何把知识组织成可回答问题的证据链”的问题，而基础设施支撑模块则负责为上述能力提供运行基础。各模块功能边界清晰，有利于后续独立维护和局部替换。")
    current = insert_captioned_image(after_module, "图 4-2 系统功能模块设计图", assets["module_design"], "在总体设计层面，为了把系统的角色入口、治理链路和服务链路对应到可实现的模块边界，本文进一步绘制了功能模块设计图。", "图 4-2 将系统划分为用户端、管理端、知识治理、文件处理、实体确认、知识图谱、问答服务和基础设施八个模块。这种划分与项目当前的前后端目录结构和控制器职责基本一致，能够为后续详细设计章节提供稳定的结构基础。")
    flow_anchor = find_paragraph(doc, "在智能问答流程中，用户提出问题后，系统需要先完成关键词检索与向量召回，再通过融合排序选取相关证据，随后构建上下文并生成回答。若问题基于上传附件产生，系统还需要维护附件与当前会话之间的关联关系，以支持连续追问。通过对关键流程的统一设计，系统能够在知识治理与知识服务之间建立稳定联系，从而形成完整的业务闭环。")
    current = insert_captioned_image(flow_anchor, "图 4-4 知识入库处理流程图", assets["ingest_flow"], "总体设计不仅需要说明模块划分，还需要说明知识从原始资料到可服务资产的关键流转链路。", "图 4-4 对知识入库流程进行了分阶段表达，体现出系统采用“资料接入、任务处理、候选抽取、人工审核、结果沉淀”的治理思路，能够较好支撑美业知识的持续积累与维护。", width_cm=12.5)
    insert_captioned_image(current, "图 4-5 用户问答服务流程图", assets["qa_flow"], "与知识入库流程相对应，用户问答服务流程则体现了系统如何把沉淀后的知识重新组织为可交互的回答能力。", "图 4-5 展示了从问题输入、检索召回、融合排序到流式返回的整体链路，说明总体设计已经把知识底座、检索服务和前端交互统一纳入一条可执行的问答服务流程。", width_cm=12.5)


def expand_intro(doc: Document) -> None:
    para = find_paragraph_contains(doc, "围绕美妆")
    insert_body_after(para, "对于美妆场景而言，知识服务对象既可能是普通消费者，也可能是品牌培训人员、门店顾问和平台运营人员。不同角色对知识的需求存在明显差异：前者更关注“能否快速获得可信解释”，后者更关注“能否持续沉淀可治理的知识资产”。这意味着系统在设计时必须兼顾前台服务效率与后台治理深度，不能只停留在静态知识展示层面。")
    para2 = find_paragraph_contains(doc, "综合型系统仍相对较少")
    insert_body_after(para2, "进一步来看，毕业设计论文并不仅要求完成系统编码，更强调将系统建设过程中的需求分析、技术选型、结构设计、接口组织和测试验证串联为完整论证。因此，本文在绪论中不仅关注课题背景，也强调该课题与知识治理、文档智能处理、图谱组织和检索增强问答之间的关系，为后文展开系统化论述建立逻辑起点。")
    para3 = find_paragraph_contains(doc, "从具体研究任务来看")
    insert_body_after(para3, "从研究主线看，本文重点围绕知识治理、文件处理、实体抽取确认、图谱构建和智能问答五个方面展开。其中，知识治理决定平台的数据质量基础，文件处理决定原始资料能否稳定进入系统，实体确认和图谱构建决定结构化知识是否可靠，智能问答则决定这些知识资产能否最终转化为用户可感知的服务能力。")


def remove_old_screenshot_block(doc: Document) -> None:
    start = "为验证系统登录入口在真实运行环境下可以正常展示，本文首先对系统登录页面进行截图取证。该页面提供管理员与普通用户两类入口，是后续所有业务流程的统一起点。"
    end = "截图展示了公告列表、发布状态与操作按钮，表明系统能够完成公告信息的查看、发布与管理。"
    for para in list(paragraphs_between(doc, start, end)):
        remove_paragraph(para)


def insert_detailed_design_showcase(doc: Document, assets: dict[str, Path]) -> None:
    after_interface = find_paragraph(doc, "这些接口共同构成了系统从前端交互到后端服务处理的主要通道。接口设计并不追求复杂，而是强调与实际业务流程的一致性：即前端页面的每一步操作都能够找到明确的后端接口响应，后端返回的数据结构又能够支撑页面的有效展示。通过这种方式，系统在详细实现层面形成了较为稳定的前后端协同机制。")
    current = build_interface_table(after_interface, doc)
    current = insert_heading2_after(current, "5.8 系统界面与关键代码展示")
    current = insert_heading3_after(current, "5.8.1 运行页面展示")
    current = insert_body_after(current, "为了使系统详细设计不仅停留在流程和接口说明层面，本文进一步给出用户端与管理端核心页面的真实运行截图。这些界面截图与前文的模块设计、接口说明和业务流程相互对应，可作为系统实现结果的直观证明。")
    for block in screenshot_blocks():
        current = insert_captioned_image(current, block.title, block.path, block.intro, block.analysis, block.width_cm)
    current = insert_heading3_after(current, "5.8.2 关键代码截图")
    current = insert_body_after(current, "除页面截图外，论文还需要适度展示能够体现系统核心能力的关键代码实现。为避免大段原始代码影响阅读，本文采用白底代码截图方式呈现前端路由、知识管理、文件处理、知识图谱与智能问答等关键实现片段。")
    for block in code_blocks(assets):
        current = insert_captioned_image(current, block.title, block.path, block.intro, block.analysis, block.width_cm)


def apply_heading_styles(doc: Document) -> None:
    for para in doc.paragraphs:
        text = para.text.strip()
        if re.match(r"^第[一二三四五六七八九十]+章", text) or text in {"结    论", "参考文献", "致    谢"}:
            style_heading_1(para)
        elif re.match(r"^\d+\.\d+\.\d+\s", text):
            style_heading_3(para)
        elif re.match(r"^\d+\.\d+\s", text):
            style_heading_2(para)
        elif re.match(r"^[图表]\s", text):
            style_caption(para)


def normalize_body_styles(doc: Document) -> None:
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or re.match(r"^(第[一二三四五六七八九十]+章|\d+\.\d+(\.\d+)?\s|[图表]\s|结\s*论|参考文献|致\s*谢)", text):
            continue
        style_body_paragraph(para)


def restructure() -> Path:
    assets = generate_assets()
    WORK_DOCX.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(SOURCE_DOCX, WORK_DOCX)
    doc = Document(WORK_DOCX)
    expand_intro(doc)
    insert_chapter2(doc)
    update_existing_structure(doc)
    insert_requirement_data_flow(doc, assets)
    insert_overall_design_assets(doc, assets)
    remove_old_screenshot_block(doc)
    insert_detailed_design_showcase(doc, assets)
    apply_heading_styles(doc)
    normalize_body_styles(doc)
    add_reference_hyperlinks(doc)
    doc.save(WORK_DOCX)
    return WORK_DOCX


def main() -> None:
    print(restructure())


if __name__ == "__main__":
    main()
