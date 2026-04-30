from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\beauty\毕业论文_最终稿_规范达标版.docx")


def style_run(run, font_name: str = "宋体", size_pt: float = 14, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    for run in paragraph.runs:
        style_run(run, "宋体", 14)


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_text(paragraph: Paragraph, text: str) -> None:
    p_pr = paragraph._p.pPr
    clear_paragraph(paragraph)
    if p_pr is not None and paragraph._p.pPr is None:
        paragraph._p.insert(0, p_pr)
    paragraph.add_run(text)
    style_body_paragraph(paragraph)


def find_exact(doc: Document, text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Paragraph not found: {text}")


def remove_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)
        paragraph._p = paragraph._element = None


def remove_figure_block(doc: Document, caption: str) -> None:
    caption_para = find_exact(doc, caption)
    paras = list(doc.paragraphs)
    idx = next(i for i, p in enumerate(paras) if p.text.strip() == caption)
    block = []
    for target in (idx - 2, idx - 1, idx, idx + 1):
        if 0 <= target < len(paras):
            block.append(paras[target])
    for para in reversed(block):
        remove_paragraph(para)


def main() -> None:
    doc = Document(DOCX_PATH)

    remove_captions = [
        "图 5-8 普通用户端首页页面截图",
        "图 5-10 普通用户端收藏页面截图",
        "图 5-11 管理员端运营总览页面截图",
        "图 5-12 管理员端运营报表页面截图",
        "图 5-15 管理员端分类树页面截图",
        "图 5-19 管理员端用户管理页面截图",
        "图 5-20 管理员端公告管理页面截图",
    ]
    for caption in remove_captions:
        remove_figure_block(doc, caption)

    replacements = {
        "为了使系统详细设计不仅停留在流程和接口说明层面，本文进一步给出用户端与管理端核心页面的真实运行截图。这些界面截图与前文的模块设计、接口说明和业务流程相互对应，可作为系统实现结果的直观证明。":
        "为了避免运行界面展示沦为页面罗列，本文仅选取最能体现系统技术主线的核心功能页面作为代表性运行截图。所保留的界面分别对应统一认证入口、检索增强问答、知识治理、文件接入、实体审核、图谱展示和任务监控等关键链路，可直接支撑详细设计章节的技术论述。",
        "图 5-9 普通用户端智能问答页面截图": "图 5-8 普通用户端智能问答页面截图",
        "图 5-9 对应智能问答模块的前端承载界面。会话列表、消息区域、快捷提问和输入控件共同说明该页面不仅负责消息展示，还负责会话状态维护、问题输入组织以及与流式问答接口的协同调用，是问答服务链路在界面层的核心载体。":
        "图 5-8 对应智能问答模块的前端承载界面。会话列表、消息区域、快捷提问和输入控件共同说明该页面不仅负责消息展示，还负责会话状态维护、问题输入组织以及与流式问答接口的协同调用，是问答服务链路在界面层的核心载体。",
        "图 5-13 管理员端知识列表页面截图": "图 5-9 管理员端知识列表页面截图",
        "图 5-13 对应知识管理模块的核心治理界面。该页面将分页列表、筛选条件和状态维护入口统一组织，说明前端页面已经与知识分页、搜索和状态更新接口形成闭环协同，是系统知识治理能力在界面层的直接体现。":
        "图 5-9 对应知识管理模块的核心治理界面。该页面将分页列表、筛选条件和状态维护入口统一组织，说明前端页面已经与知识分页、搜索和状态更新接口形成闭环协同，是系统知识治理能力在界面层的直接体现。",
        "图 5-14 管理员端文件上传页面截图": "图 5-10 管理员端文件上传页面截图",
        "图 5-14 对应文件接入与任务创建页面。该页面通过上传控件、知识信息表单和提交入口，将资料接入动作显式转换为后端任务登记和处理链路触发点，体现了文件处理模块在界面层对异步任务机制的承载方式。":
        "图 5-10 对应文件接入与任务创建页面。该页面通过上传控件、知识信息表单和提交入口，将资料接入动作显式转换为后端任务登记和处理链路触发点，体现了文件处理模块在界面层对异步任务机制的承载方式。",
        "图 5-16 管理员端实体确认页面截图": "图 5-11 管理员端实体确认页面截图",
        "图 5-16 对应实体审核页面的真实运行状态。待审核候选、统计信息与审核动作区共同说明系统已将“候选生成 - 人工确认 - 结果回写”的治理机制落实到界面层，从而控制图谱数据质量并降低自动抽取误差传播风险。":
        "图 5-11 对应实体审核页面的真实运行状态。待审核候选、统计信息与审核动作区共同说明系统已将“候选生成 - 人工确认 - 结果回写”的治理机制落实到界面层，从而控制图谱数据质量并降低自动抽取误差传播风险。",
        "图 5-17 管理员端知识图谱页面截图": "图 5-12 管理员端知识图谱页面截图",
        "图 5-17 表明知识图谱页面已经能够承载图结构查询结果的可视化展示。该页面通过节点关系展示和查询入口，把结构化知识组织结果转化为可观察的图形化界面，是图谱服务能力对管理端开放的直接表现。":
        "图 5-12 表明知识图谱页面已经能够承载图结构查询结果的可视化展示。该页面通过节点关系展示和查询入口，把结构化知识组织结果转化为可观察的图形化界面，是图谱服务能力对管理端开放的直接表现。",
        "图 5-18 管理员端任务监控页面截图": "图 5-13 管理员端任务监控页面截图",
        "图 5-18 对应任务监控页面。任务状态、时间信息和重试操作共同说明系统在异步处理链路中已经建立起任务可观测、失败可重试和状态可追踪的机制，这对于多媒体资料解析场景的稳定运行尤为关键。":
        "图 5-13 对应任务监控页面。任务状态、时间信息和重试操作共同说明系统在异步处理链路中已经建立起任务可观测、失败可重试和状态可追踪的机制，这对于多媒体资料解析场景的稳定运行尤为关键。",
        "图 5-21 选取前端路由定义与守卫代码，用于说明系统如何在页面层区分普通用户与管理员入口。": "图 5-14 选取前端路由定义与守卫代码，用于说明系统如何在页面层区分普通用户与管理员入口。",
        "图 5-21 前端统一路由与角色鉴权代码截图": "图 5-14 前端统一路由与角色鉴权代码截图",
        "图 5-22 展示知识管理控制器的主要实现，用于说明知识分页、详情、保存、更新和搜索能力的接口组织方式。": "图 5-15 展示知识管理控制器的主要实现，用于说明知识分页、详情、保存、更新和搜索能力的接口组织方式。",
        "图 5-22 知识管理接口控制器代码截图": "图 5-15 知识管理接口控制器代码截图",
        "图 5-23 选取文件处理控制器代码，展示文件上传、任务查询、重试和文件打开等关键接口实现。": "图 5-16 选取文件处理控制器代码，展示文件上传、任务查询、重试和文件打开等关键接口实现。",
        "图 5-23 文件上传与任务追踪接口代码截图": "图 5-16 文件上传与任务追踪接口代码截图",
        "图 5-24 选取图谱控制器代码，用于展示候选审核、图谱查询、路径分析和证据查看接口的组织方式。": "图 5-17 选取图谱控制器代码，用于展示候选审核、图谱查询、路径分析和证据查看接口的组织方式。",
        "图 5-24 知识图谱接口控制器代码截图": "图 5-17 知识图谱接口控制器代码截图",
        "图 5-25 展示智能问答控制器中流式问答、附件摘要、附件追问和会话管理相关代码。": "图 5-18 展示智能问答控制器中流式问答、附件摘要、附件追问和会话管理相关代码。",
        "图 5-25 智能问答流式接口代码截图": "图 5-18 智能问答流式接口代码截图",
    }

    for para in doc.paragraphs:
        text = para.text.strip()
        if text in replacements:
            set_text(para, replacements[text])

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
