from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\beauty")
DOCX_PATH = ROOT / "毕业论文_最终稿_规范达标版.docx"
ASSET_DIR = ROOT / "final_assets" / "thesis_restructure"
BACKEND_DIR = ROOT / "beauty-knowledge-backend"
FRONTEND_DIR = ROOT / "beauty-knowledge-frontend"
FONT_BOLD = Path(r"C:\Windows\Fonts\simhei.ttf")
FONT_MONO = Path(r"C:\Windows\Fonts\consola.ttf")


def style_run(run, font_name: str = "宋体", size_pt: float = 14, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    for run in paragraph.runs:
        style_run(run, "宋体", 14)


def style_caption(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    for run in paragraph.runs:
        style_run(run, "宋体", 14)


def style_image_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.0


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def find_exact(doc: Document, text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Paragraph not found: {text}")


def remove_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)
        paragraph._p = paragraph._element = None


def get_font(path: Path, size: int):
    return ImageFont.truetype(str(path), size=size)


def read_lines(path: Path, start: int, end: int) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    picked = []
    for idx in range(start - 1, min(end, len(lines))):
        picked.append(f"{idx + 1:>3}  {lines[idx].rstrip()}")
    return "\n".join(picked)


def create_code_image(output: Path, title: str, code: str) -> None:
    title_font = get_font(FONT_BOLD, 26)
    code_font = get_font(FONT_MONO, 22)
    line_height = 34
    padding = 32
    title_h = 70
    lines = code.splitlines() or [""]
    width = 1700
    height = title_h + padding * 2 + len(lines) * line_height + 20
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, title_h), fill="#f5f5f5", outline="black", width=2)
    draw.text((28, 20), title, fill="black", font=title_font)
    y = title_h + padding
    for line in lines:
        draw.text((padding, y), line.replace("\t", "    "), fill="black", font=code_font)
        y += line_height
    draw.rectangle((1, 1, width - 2, height - 2), outline="black", width=2)
    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)


def generate_core_images() -> dict[str, Path]:
    assets = {
        "router": ASSET_DIR / "fig_5_14_code_router_core.png",
        "knowledge": ASSET_DIR / "fig_5_15_code_knowledge_core.png",
        "file": ASSET_DIR / "fig_5_16_code_file_core.png",
        "kg": ASSET_DIR / "fig_5_17_code_kg_core.png",
        "chat": ASSET_DIR / "fig_5_18_code_chat_core.png",
    }
    create_code_image(
        assets["router"],
        "前端统一路由与角色鉴权核心代码",
        read_lines(FRONTEND_DIR / "src" / "router" / "index.ts", 16, 61),
    )
    create_code_image(
        assets["knowledge"],
        "知识管理接口控制器核心代码",
        read_lines(BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "cms" / "controller" / "KnowledgeController.java", 22, 67),
    )
    create_code_image(
        assets["file"],
        "文件上传与任务追踪核心代码",
        read_lines(BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "cms" / "controller" / "FileController.java", 30, 86),
    )
    create_code_image(
        assets["kg"],
        "知识图谱治理与查询核心代码",
        read_lines(BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "kg" / "controller" / "KgController.java", 33, 111),
    )
    create_code_image(
        assets["chat"],
        "智能问答流式输出核心代码",
        read_lines(BACKEND_DIR / "src" / "main" / "java" / "com" / "beauty" / "knowledge" / "module" / "rag" / "controller" / "ChatController.java", 31, 118),
    )
    return assets


def remove_old_code_block(doc: Document) -> None:
    start_text = "图 5-14 选取前端路由定义与守卫代码，用于说明系统如何在页面层区分普通用户与管理员入口。"
    end_text = "该段代码在整体架构中承担问答服务对外输出职责。通过 SSE 流式返回、附件摘要、附件追问和会话管理接口的组合，系统把检索增强链路、上下文维护和交互式问答统一组织在同一控制器中，解决了复杂问答场景下结果实时输出与会话持续管理的实现问题。"
    collecting = False
    to_remove = []
    for para in list(doc.paragraphs):
        txt = para.text.strip()
        if txt == start_text:
            collecting = True
        if collecting:
            to_remove.append(para)
        if txt == end_text and collecting:
            break
    for para in reversed(to_remove):
        remove_paragraph(para)


def insert_code_block(after: Paragraph, intro: str, caption: str, image_path: Path, analysis: str) -> Paragraph:
    intro_p = insert_paragraph_after(after, intro)
    style_body_paragraph(intro_p)
    img_p = insert_paragraph_after(intro_p)
    img_p.add_run().add_picture(str(image_path), width=Cm(13.6))
    style_image_paragraph(img_p)
    cap_p = insert_paragraph_after(img_p, caption)
    style_caption(cap_p)
    ana_p = insert_paragraph_after(cap_p, analysis)
    style_body_paragraph(ana_p)
    return ana_p


def main() -> None:
    assets = generate_core_images()
    doc = Document(DOCX_PATH)
    remove_old_code_block(doc)
    anchor = find_exact(doc, "除页面截图外，论文还需要适度展示能够体现系统核心能力的关键代码实现。为避免大段原始代码影响阅读，本文采用白底代码截图方式呈现前端路由、知识管理、文件处理、知识图谱与智能问答等关键实现片段。")
    set_text = anchor  # keep anchor unchanged, insert after it
    current = anchor
    blocks = [
        (
            "图 5-14 选取前端路由定义与守卫中的关键片段，重点展示用户端与管理端路由划分以及登录态校验逻辑。",
            "图 5-14 前端统一路由与角色鉴权核心代码截图",
            assets["router"],
            "该段核心代码集中体现了前端入口控制策略：一方面通过 `/user` 与 `/admin` 两类路由树完成角色分流，另一方面通过 `beforeEach` 守卫统一处理登录态校验和越权跳转，从而避免无关辅助代码干扰论文阅读。",
        ),
        (
            "图 5-15 仅保留知识控制器中与分页查询、详情读取和状态维护直接相关的核心接口实现，以突出资源模型与权限边界设计。",
            "图 5-15 知识管理接口核心代码截图",
            assets["knowledge"],
            "该片段聚焦知识资源的核心读写接口，能够直接说明知识管理模块如何通过统一资源路径、分页查询和 `PreAuthorize` 约束实现“同一资源模型下的差异化访问控制”。",
        ),
        (
            "图 5-16 选取文件控制器中最能体现异步处理链路的核心方法，突出上传入口、任务查询与重试控制机制。",
            "图 5-16 文件上传与任务追踪核心代码截图",
            assets["file"],
            "该截图不再展示全部辅助接口，而是保留上传、任务查询和重试等关键片段，从而更直接地说明文件处理模块如何把资料接入组织为可追踪、可恢复的异步任务链路。",
        ),
        (
            "图 5-17 仅展示图谱控制器中与候选审核、图谱查询和证据访问直接相关的核心实现，用于突出图谱模块的双重职责。",
            "图 5-17 知识图谱治理与查询核心代码截图",
            assets["kg"],
            "该核心代码片段能够更集中地反映图谱模块的两类关键能力：其一是通过审核接口控制结构化结果写入质量，其二是通过图谱、路径和证据接口对外提供查询服务，从而避免非核心辅助逻辑分散论述重点。",
        ),
        (
            "图 5-18 保留问答控制器中最能体现检索增强问答输出特征的核心实现，包括流式问答、附件摘要和会话组织相关方法。",
            "图 5-18 智能问答流式输出核心代码截图",
            assets["chat"],
            "该片段集中展示了问答模块如何通过 SSE 输出、附件摘要和会话管理组织服务入口，更适合作为论文中的核心代码证据来说明复杂问答场景下的交互式输出机制。",
        ),
    ]
    for intro, caption, image_path, analysis in blocks:
        current = insert_code_block(current, intro, caption, image_path, analysis)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
