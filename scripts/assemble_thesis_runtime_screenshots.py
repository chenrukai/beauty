from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(r"D:\beauty")
SOURCE = ROOT / "毕业论文_最终稿_规范达标版.docx"
BASELINE = ROOT / "_tmp_plan_doc_copy.docx"
WORKING_COPY = ROOT / "_tmp_runtime_screenshot_build.docx"
MANIFEST_PATH = ROOT / "scripts" / "thesis_page_screenshot_manifest.json"
INSERT_AFTER_TEXT = "图 5-2 兼容性测试结果图"


def insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def set_east_asia_font(paragraph: Paragraph, font_name: str, font_size: Pt) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = font_size
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    set_east_asia_font(paragraph, "宋体", Pt(14))


def style_caption_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    set_east_asia_font(paragraph, "宋体", Pt(14))


def style_image_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Pt(0)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(6)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = 1


def find_paragraph(doc: Document, exact_text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == exact_text:
            return para
    raise ValueError(f"Paragraph not found: {exact_text}")


def load_manifest() -> list[dict]:
    return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))


def build_document() -> None:
    if not BASELINE.exists():
        raise FileNotFoundError(f"Baseline copy not found: {BASELINE}")

    shutil.copyfile(BASELINE, WORKING_COPY)
    doc = Document(WORKING_COPY)
    manifest = sorted(load_manifest(), key=lambda item: int(item["figure_no"].split("-")[1]))

    insert_after = find_paragraph(doc, INSERT_AFTER_TEXT)

    for item in manifest:
        image_path = Path(item["image_path"])
        if not image_path.exists():
            raise FileNotFoundError(f"Screenshot not found: {image_path}")

        intro_para = insert_paragraph_after(insert_after, item["intro"])
        style_body_paragraph(intro_para)

        image_para = insert_paragraph_after(intro_para)
        image_run = image_para.add_run()
        image_run.add_picture(str(image_path), width=Cm(13.6))
        style_image_paragraph(image_para)

        caption_para = insert_paragraph_after(image_para, f"图 {item['figure_no']} {item['title']}")
        style_caption_paragraph(caption_para)

        analysis_para = insert_paragraph_after(caption_para, item["analysis"])
        style_body_paragraph(analysis_para)

        insert_after = analysis_para

    doc.save(SOURCE)


if __name__ == "__main__":
    build_document()
