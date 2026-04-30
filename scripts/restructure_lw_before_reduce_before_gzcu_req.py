from copy import deepcopy
from pathlib import Path
import re

from docx import Document


DOC_PATH = Path(r"D:\beauty\lw_before_reduce_before_gzcu_req.docx")
BACKUP_PATH = Path(r"D:\beauty\lw_before_reduce_before_gzcu_req_before_restructure.docx")


def insert_copies_before(target_para, source_paras):
    for src in source_paras:
        target_para._p.addprevious(deepcopy(src._p))


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)
    paragraph._p = paragraph._element = None


def replace_text_global(text: str) -> str:
    text = text.replace("第三章 系统需求与分析", "第二章 系统需求分析")
    text = text.replace("第四章 系统总体设计", "第三章 系统概要设计")
    text = text.replace("第五章 系统详细设计与实现", "第四章 系统详细设计与实现")
    text = text.replace("第六章 系统测试", "第五章 系统测试")

    text = text.replace(
        "全文结构安排如下：第一章为绪论，主要阐述研究背景、研究意义、国内外研究现状及论文结构；第二章介绍系统相关技术与理论基础；第三章分析系统需求与数据流；第四章介绍系统总体设计与关键流程；第五章详细说明系统各模块的设计实现、运行界面与关键代码；第六章为系统测试；第七章为系统部署与运行验证。",
        "全文结构安排如下：第一章为绪论，主要阐述研究背景、研究意义、国内外研究现状、技术路线与论文结构；第二章分析系统需求与数据流；第三章介绍系统概要设计与关键流程；第四章详细说明系统各模块的设计实现、运行界面与关键代码；第五章为系统测试，并补充部署与运行验证内容。"
    )

    # Heading renumbering
    text = re.sub(r"^3\.(\d+(?:\.\d+)*)", lambda m: "2." + m.group(1), text)
    text = re.sub(r"^4\.(\d+(?:\.\d+)*)", lambda m: "3." + m.group(1), text)
    text = re.sub(r"^5\.(\d+(?:\.\d+)*)", lambda m: "4." + m.group(1), text)
    text = re.sub(r"^6\.1\b", "5.1", text)
    text = re.sub(r"^6\.2\b", "5.2", text)
    text = re.sub(r"^6\.3\b", "5.3", text)
    text = re.sub(r"^6\.4\b", "5.4", text)
    text = re.sub(r"^6\.5\b", "5.5", text)
    text = re.sub(r"^6\.6\b", "5.7", text)
    text = re.sub(r"^7\.1\b", "5.6", text)
    text = re.sub(r"^7\.2\b", "5.6.1", text)
    text = re.sub(r"^7\.3\b", "5.6.2", text)

    # Temporary placeholders for figure/table renumbering
    replacements = [
        ("图 3-", "__FIG3__"),
        ("图3-", "__FIG3N__"),
        ("表 3-", "__TAB3__"),
        ("表3-", "__TAB3N__"),
        ("图 4-", "__FIG4__"),
        ("图4-", "__FIG4N__"),
        ("表 4-", "__TAB4__"),
        ("表4-", "__TAB4N__"),
        ("图 5-", "__FIG5__"),
        ("图5-", "__FIG5N__"),
        ("表 5-", "__TAB5__"),
        ("表5-", "__TAB5N__"),
        ("图 6-", "__FIG6__"),
        ("图6-", "__FIG6N__"),
        ("表 6-", "__TAB6__"),
        ("表6-", "__TAB6N__"),
    ]
    for old, token in replacements:
        text = text.replace(old, token)

    final_map = {
        "__FIG3__": "图 2-",
        "__FIG3N__": "图2-",
        "__TAB3__": "表 2-",
        "__TAB3N__": "表2-",
        "__FIG4__": "图 3-",
        "__FIG4N__": "图3-",
        "__TAB4__": "表 3-",
        "__TAB4N__": "表3-",
        "__FIG5__": "图 4-",
        "__FIG5N__": "图4-",
        "__TAB5__": "表 4-",
        "__TAB5N__": "表4-",
        "__FIG6__": "图 5-",
        "__FIG6N__": "图5-",
        "__TAB6__": "表 5-",
        "__TAB6N__": "表5-",
    }
    for token, new in final_map.items():
        text = text.replace(token, new)

    return text


def main():
    if not BACKUP_PATH.exists():
        BACKUP_PATH.write_bytes(DOC_PATH.read_bytes())

    doc = Document(str(DOC_PATH))
    paras = doc.paragraphs

    # Store references before structural edits
    target_ch1_summary = paras[67]
    target_ch3_fig = paras[132]
    target_ch3_er_fig = paras[141]
    target_ch4_qa_fig = paras[180]

    tech_stack_paras = paras[71:73]
    media_paras = paras[74:76]
    kg_paras = paras[77:79]
    rag_paras = paras[80:82]

    insert_copies_before(target_ch1_summary, tech_stack_paras)
    insert_copies_before(target_ch3_fig, media_paras)
    insert_copies_before(target_ch3_er_fig, kg_paras)
    insert_copies_before(target_ch4_qa_fig, rag_paras)

    # Merge old chapter 7 summary into chapter 5 summary paragraph
    if paras[260].text.strip() and paras[276].text.strip():
        paras[260].text = paras[260].text.rstrip("。") + "。此外，" + paras[276].text.lstrip("本章").lstrip()

    # Rename merged heading
    paras[262].text = "5.6 系统部署与运行验证"
    paras[267].text = "5.6.1 运行链路验证"
    paras[272].text = "5.6.2 工程化特点分析"

    # Delete old chapter 2 and old chapter 7 shell paragraphs
    to_delete = list(paras[69:83]) + [paras[261], paras[275], paras[276]]
    for para in reversed(to_delete):
        delete_paragraph(para)

    # Global text updates after structure changes
    for para in doc.paragraphs:
        text = para.text
        if not text:
            continue
        new_text = replace_text_global(text)
        if new_text != text:
            para.text = new_text

    doc.save(str(DOC_PATH))


if __name__ == "__main__":
    main()
