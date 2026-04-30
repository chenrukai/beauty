from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\beauty\毕业论文_最终稿_规范达标版.docx")


def style_run(run, font_name: str = "宋体", size_pt: float = 10.5, superscript: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.superscript = superscript
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def rebuild_citation_runs(paragraph: Paragraph) -> None:
    text = paragraph.text
    if not re.search(r"\[(\d+)\]", text):
        return
    p_pr = paragraph._p.pPr
    clear_paragraph(paragraph)
    if p_pr is not None and paragraph._p.pPr is None:
        paragraph._p.insert(0, p_pr)
    cursor = 0
    for match in re.finditer(r"\[(\d+)\]", text):
        normal = text[cursor:match.start()]
        if normal:
            run = paragraph.add_run(normal)
            style_run(run, "宋体", 10.5, superscript=False)
        run = paragraph.add_run(f"[{match.group(1)}]")
        style_run(run, "宋体", 10.5, superscript=True)
        cursor = match.end()
    tail = text[cursor:]
    if tail:
        run = paragraph.add_run(tail)
        style_run(run, "宋体", 10.5, superscript=False)
    style_body_paragraph(paragraph)


def main() -> None:
    doc = Document(DOCX_PATH)
    body_started = False
    skip_summary = False
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt == "第一章 绪论":
            body_started = True
        if txt in {"结    论", "结论"}:
            break
        if not body_started or not txt:
            continue
        if txt.endswith("本章小结"):
            skip_summary = True
            continue
        if skip_summary:
            if re.match(r"^(第[一二三四五六七八九十]+章|\d+\.\d+(\.\d+)?\s)", txt):
                skip_summary = False
            else:
                continue
        if re.search(r"\[(\d+)\]", txt):
            rebuild_citation_runs(para)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
