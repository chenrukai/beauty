from pathlib import Path

from docx import Document


DOC_PATH = Path(r"D:\beauty\lw_before_reduce_before_gzcu_req.docx")


def delete_paragraph(paragraph):
    p = paragraph._element
    parent = p.getparent()
    if parent is not None:
        parent.remove(p)
    paragraph._p = paragraph._element = None


def main():
    doc = Document(str(DOC_PATH))
    paras = doc.paragraphs

    paras[66].text = "全文结构安排如下：第一章为绪论，主要阐述研究背景、研究意义、国内外研究现状、技术路线与论文结构；第二章分析系统需求与数据流；第三章介绍系统概要设计与关键流程；第四章详细说明系统各模块的设计实现、运行界面与关键代码；第五章为系统测试，并补充部署与运行验证内容；最后对全文进行总结，并给出后续改进方向。"
    paras[70].text = "本章围绕课题背景、研究意义、国内外研究现状、技术路线以及论文整体结构进行了系统阐述，明确了本文所关注的关键技术问题与研究主线，并从总体上界定了多源资料接入、知识治理闭环、结构化知识组织和检索增强生成问答之间的内在联系，为后续需求分析、概要设计和详细实现提供了论证基础。"
    delete_paragraph(paras[71])

    paras = doc.paragraphs
    paras[149].text = "第四章 系统详细设计与实现"
    paras[150].text = "在总体设计基础上，本章进一步从模块职责、接口组织、处理流程和前后端协同关系四个方面说明系统的详细设计与实现，如图 4-1 所示。该协作图用于概括前端页面、后端接口、数据存储和智能服务之间的总体配合关系，使读者在进入各模块细节前，能够先对系统实现链路形成整体认识。"
    paras[153].text = "4.1 前端交互模块设计与实现"
    paras[157].text = "4.2 知识管理模块设计与实现"
    paras[161].text = "4.3 文件处理模块设计与实现"
    paras[165].text = "4.4 实体抽取与知识图谱模块设计与实现"
    paras[169].text = "4.5 智能问答模块设计与实现"
    paras[175].text = "4.6 接口设计说明"
    paras[180].text = "4.7 系统界面与关键代码展示"
    paras[181].text = "4.7.1 运行页面展示"
    paras[204].text = "4.7.2 关键代码截图"
    paras[221].text = "4.8 本章小结"

    paras[255].text = "5.6 系统部署与运行验证"
    paras[260].text = "5.6.1 运行链路验证"
    paras[265].text = "5.6.2 工程化特点分析"

    doc.save(str(DOC_PATH))


if __name__ == "__main__":
    main()
