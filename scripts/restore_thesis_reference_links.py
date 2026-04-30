from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\beauty\毕业论文_最终稿_规范达标版.docx")


def style_run(run, font_name: str = "宋体", size_pt: float = 10.5) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_body_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(0.74)
    fmt.line_spacing = Pt(28)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def add_bookmark(paragraph: Paragraph, bookmark_name: str, bookmark_id: int) -> None:
    if any(
        child.tag == qn("w:bookmarkStart") and child.get(qn("w:name")) == bookmark_name
        for child in paragraph._p.iterchildren()
    ):
        return
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def append_internal_hyperlink(paragraph: Paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    vert = OxmlElement("w:vertAlign")
    vert.set(qn("w:val"), "superscript")
    rpr.append(vert)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    rpr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "21")
    rpr.append(size_cs)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def rebuild_citations(paragraph: Paragraph) -> None:
    text = paragraph.text
    if not re.search(r"\[(\d+)\]", text):
        return
    p_pr = paragraph._p.pPr
    clear_paragraph(paragraph)
    if p_pr is not None and paragraph._p.pPr is None:
        paragraph._p.insert(0, p_pr)
    cursor = 0
    for match in re.finditer(r"\[(\d+)\]", text):
        normal = text[cursor:match.start()]
        if normal:
            run = paragraph.add_run(normal)
            style_run(run)
        ref_no = match.group(1)
        append_internal_hyperlink(paragraph, f"[{ref_no}]", f"gzcu_ref_{ref_no}")
        cursor = match.end()
    tail = text[cursor:]
    if tail:
        run = paragraph.add_run(tail)
        style_run(run)
    style_body_paragraph(paragraph)


def main() -> None:
    doc = Document(DOCX_PATH)
    bookmark_id = 900
    in_refs = False
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt == "参考文献":
            in_refs = True
            continue
        if in_refs:
            if txt in {"致    谢", "致谢"}:
                break
            match = re.match(r"^\[(\d+)\]", txt)
            if match:
                add_bookmark(para, f"gzcu_ref_{match.group(1)}", bookmark_id)
                bookmark_id += 1

    body_started = False
    skip_summary = False
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt == "第一章 绪论":
            body_started = True
        if txt in {"结    论", "结论"}:
            break
        if not body_started or not txt:
            continue
        if txt.endswith("本章小结"):
            skip_summary = True
            continue
        if skip_summary:
            if re.match(r"^(第[一二三四五六七八九十]+章|\d+\.\d+(\.\d+)?\s)", txt):
                skip_summary = False
            else:
                continue
        if re.search(r"\[(\d+)\]", txt):
            rebuild_citations(para)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
