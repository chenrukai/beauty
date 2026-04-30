from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document


DOCX_PATH = Path(r"D:\beauty\lw.docx")
BACKUP_PATH = Path(r"D:\beauty\lw_before_final_consistency_fix.docx")


PARA_REPLACEMENTS = {
    36: "针对美妆领域知识来源异构、结构化组织程度不足、资料复用效率较低以及专业问答缺乏稳定证据支撑等问题，本文设计并实现了一套面向知识治理与知识服务协同的美业多媒体知识库系统。系统以多源资料接入、知识治理闭环、结构化知识组织和检索增强生成问答为主线，在前端采用 Vue 3、Vite、Pinia 与 Element Plus 构建用户端和管理端界面，在后端基于 Spring Boot、Spring Security、MyBatis-Plus 与 WebFlux 实现认证授权、知识管理、图谱服务和流式问答能力。数据层综合使用 MySQL、Redis、RabbitMQ、MinIO 与 Milvus，分别支撑业务实体存储、缓存访问、异步任务调度、原始文件留存与向量检索；OCR 服务用于处理文档、图片与扫描资料。系统通过关键词检索、向量召回、融合排序与证据回填构建检索增强生成链路，并结合实体抽取确认与图谱构建机制提升知识组织质量与回答可解释性。论文基于真实项目代码、数据库结构、接口控制器和部署配置，对系统需求分析、总体设计、详细实现、工程测试与运行验证进行了系统论述。研究结果表明，该系统能够较好地完成美妆知识的汇聚、治理、检索、问答和图谱分析任务，具有一定的工程实现价值与应用参考意义。",
    39: "This thesis designs and implements a CMS-based multimedia beauty knowledge base construction and presentation system for the problems of scattered information sources, weak structured organization and inefficient professional retrieval in the beauty domain. The system integrates knowledge management, file processing, entity confirmation, knowledge graph visualization, hybrid retrieval question answering, statistical reporting and containerized deployment. The frontend is built with Vue 3, Vite, Pinia and Element Plus, while the backend is developed with Spring Boot, Spring Security, MyBatis-Plus and WebFlux. MySQL, Redis, RabbitMQ, MinIO and Milvus are jointly used for storage, cache, asynchronous processing, object storage and vector retrieval. OCR services are used for document and image parsing. A hybrid retrieval pipeline combining keyword search, vector recall, rank fusion and evidence backfilling is introduced to improve both relevance and interpretability. Based on the real project codebase, database schema, controllers and deployment configuration, this thesis presents the requirement analysis, architecture design, detailed implementation, testing and deployment verification of the system. The study shows that the system can effectively support beauty knowledge aggregation, management, retrieval, question answering and graph analysis.",
    104: "多媒体资料解析技术是系统实现知识接入的关键基础。与直接录入结构化知识不同，美妆领域存在大量说明书、培训资料、图片和扫描页等原始载体，这些资料通常缺乏统一格式，难以直接进入知识服务链路。因此，本系统在接入层引入 MinIO 对象存储与 OCR 文本识别能力，将原始文件逐步转换为可清洗、可切片、可抽取的文本中间结果。文档智能处理研究表明，对非结构化资料进行识别、解析和结构化组织，是构建行业知识底座的重要前提[9][10]。",
    117: "从技术可行性角度看，本系统采用的前端与后端技术栈较为成熟，具备良好的开发生态和文档支持。前端基于 Vue 3 构建页面交互，后端基于 Spring Boot 生态实现业务逻辑，结合 MySQL、Redis、RabbitMQ、MinIO、Milvus 与 OCR 等组件，能够满足系统在知识存储、缓存、异步处理、对象存储、向量检索和文件智能解析方面的实际需求。因此，系统具备较好的技术实现基础。",
    125: "在该架构中，用户端与管理端共享统一的服务后端，但在页面入口和业务职责上保持相互独立。整体结构按照访问层、接口/应用层、业务服务层和数据与 AI 支撑层进行组织，其中认证授权、知识管理、文件处理、图谱服务、问答服务与统计分析共同构成核心业务能力，MySQL、Redis、RabbitMQ、MinIO、Milvus 与 OCR 等组件提供底层支撑。系统整体结构如图 3-1 所示。",
    191: "系统接收文件上传请求后，首先完成文件对象存储和任务登记。原始文件被写入对象存储服务，对应的元数据则记录在文件表和任务表中。随后，系统根据文件类型选择不同处理链路：对于图片和扫描文档，调用 OCR 服务提取文字信息；对于文本型文档，则直接进行正文解析与清洗。处理完成后，文本结果进入清洗、切片和后续实体抽取流程。文件处理模块的核心流程如图 4-4 所示。",
    285: "测试环境基于当前项目实际开发配置展开，包括前端页面、Spring Boot 后端服务以及 MySQL、Redis、RabbitMQ、MinIO、Milvus 与 OCR 等配套基础组件。测试数据覆盖知识、分类、文件、任务、会话与图谱等核心业务对象，并结合文本和图片附件样本组织功能与流程验证；兼容性测试分别在 Chrome 与 Edge 稳定版下完成。性能记录则围绕 5 类关键操作展开，包括登录请求、知识列表、知识详情、问答首响和文本/图片任务处理，并以同一环境下的实际运行结果作为统计依据。表 5-1 给出了测试环境的主要配置。",
    294: "除功能正确性外，系统还需要从响应性能与运行兼容性两个方面验证其工程可用性。性能测试选取登录请求、知识列表、知识详情、问答首响和文本/图片任务处理 5 类关键操作作为观测对象，以覆盖同步接口响应与异步处理任务两类典型场景；兼容性测试则关注前端页面在主流桌面端浏览器中的显示一致性、交互完整性和入口可达性。",
    295: "在兼容性层面，系统采用前后端分离模式，前端通过浏览器访问，后端通过统一接口提供服务。因此，兼容性测试不仅要验证页面是否能够正确渲染，还要验证不同浏览器环境下的登录、知识浏览、后台治理和问答交互是否保持一致，从而确认系统具备基本跨环境运行能力。本次验证重点覆盖 Chrome 与 Edge 两类主流浏览器，并围绕登录、知识检索、文件上传、图谱查询和智能问答等核心场景记录测试结果。",
    296: "图 5-1 展示了关键业务接口与处理任务的性能测试结果，可用于观察同步访问请求与异步任务处理的耗时分布；图 5-2 给出了系统在 Chrome 与 Edge 稳定版下的兼容性测试矩阵，用于说明主要功能在不同浏览器环境中的可用性与一致性。两图共同从响应特征和跨浏览器运行效果两个角度支撑系统工程可用性分析。",
    306: "综合功能测试、性能测试、兼容性测试和安全性测试结果可以看出，系统已经形成了从知识治理到知识服务的完整工程闭环。知识管理模块能够稳定支撑知识录入、更新和检索；文件处理模块能够完成上传、任务跟踪和重试；图谱模块能够提供结构化关系查询；智能问答模块能够基于检索结果返回流式回答与来源信息。这些结果说明，前文提出的关键设计目标在实现层已获得较为充分的验证。",
    307: "从测试链路层面看，系统已经建立起“知识录入 - 文件处理 - 结构化组织 - 服务输出”的完整闭环。功能测试验证了主要业务流程的正确性，性能记录显示登录请求、知识列表与知识详情响应时间基本控制在 1 s 左右，问答首响约为 2~3 s，文本/图片任务处理平均耗时处于十几秒量级；兼容性测试则证明系统在 Chrome 与 Edge 稳定版下具备较好的功能一致性。因此，测试结果能够为前文的总体设计与详细实现论述提供直接支撑。",
    308: "与此同时，测试结果也揭示了系统仍存在进一步优化空间。例如，在更大规模知识数据、更高并发访问和更复杂附件类型场景下，仍可继续补充压力测试、异常链路测试和问答质量评估。这些内容超出了当前毕业设计的最小验证范围，但可以作为系统工程化演进的重要方向。",
    309: "总体而言，本文完成了一个面向美妆领域知识服务场景的系统化工程实现，并通过需求分析、总体设计、详细实现、测试验证和运行分析形成了较完整的论文论证链路。后续工作仍可围绕更大规模知识数据、更复杂图谱关系组织、更细粒度问答质量评估以及更充分的量化性能验证继续展开，以进一步提升系统的工程完备性与应用推广价值。",
    311: "为保证系统运行环境的统一性与可复现性，本系统采用容器化方式管理主要基础服务。通过 `docker-compose.yml` 对 MySQL、Redis、RabbitMQ、MinIO、etcd、Milvus 和 OCR AI 等组件进行集中编排，可以在较短时间内完成开发与演示环境搭建。前端与后端应用分别承担页面访问与业务服务职责，在统一基础设施支撑下协同运行。",
}


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new_text


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(str(DOCX_PATH))

    for idx, text in PARA_REPLACEMENTS.items():
        replace_paragraph_text(doc.paragraphs[idx], text)

    # Update test environment table to remove audio-related service.
    test_env_table = doc.tables[6]
    test_env_table.rows[5].cells[2].text = "Milvus、OCR 服务"

    doc.save(str(DOCX_PATH))
    print(str(DOCX_PATH))


if __name__ == "__main__":
    main()
